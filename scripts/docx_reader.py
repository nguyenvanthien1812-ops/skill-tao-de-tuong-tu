#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Module Đọc và Trích Xuất Nội Dung Đề Thi Từ Nguồn Đầu Vào Đa Dạng

Tính năng:
1. Đọc DOCX có MathType OLE (Equation.DSMT4) → trích xuất LaTeX
   API: https://mathtype-latex-api.onrender.com/api/convert-docx
   → Giáo viên có đề cũ làm bằng MathType → trích về LaTeX để AI tạo đề mới tương tự.

2. OCR ảnh chụp đề thi (PNG/JPG/JPEG) → nhận diện công thức → LaTeX
   API: https://mathtype-latex-api.onrender.com/v1/img2latex (Mistral AI Vision)
   → Giáo viên chụp ảnh đề thi bất kỳ → AI phân tích → tạo đề mới tương tự.

3. Đọc DOCX không MathType → trích xuất văn bản + LaTeX thuần túy (offline).
   Sử dụng python-docx.

4. Đọc PDF thông qua PyMuPDF (fitz) → trích xuất text.
"""

import os
import sys
import uuid
import base64
import urllib.request
import urllib.parse
import json
from pathlib import Path
from typing import Optional

# ── Dependencies ────────────────────────────────────────────────────────────
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

# ── API Endpoints ────────────────────────────────────────────────────────────
MATHTYPE_READER_API   = "https://mathtype-latex-api.onrender.com"
ENDPOINT_DOCX_READER  = f"{MATHTYPE_READER_API}/api/convert-docx"  # MathType OLE → LaTeX
ENDPOINT_IMG2LATEX    = f"{MATHTYPE_READER_API}/v1/img2latex"        # Ảnh → LaTeX (Mistral AI)

# ── License helpers (tái sử dụng từ docx_math_builder) ──────────────────────
def _get_license_token() -> str:
    try:
        lf = Path(__file__).parent.parent / "license.key"
        return lf.read_text(encoding="utf-8").strip() if lf.exists() else ""
    except Exception:
        return ""

def _require_license():
    try:
        _d = os.path.dirname(os.path.abspath(__file__))
        if _d not in sys.path:
            sys.path.insert(0, _d)
        from license_manager import check_current_license
        ok, payload, err = check_current_license()
        if not ok:
            raise PermissionError(
                "\n╔══════════════════════════════════════════════════════════╗\n"
                "║     ⛔  SKILL CHƯA ĐƯỢC KÍCH HOẠT BẢN QUYỀN!           ║\n"
                "╠══════════════════════════════════════════════════════════╣\n"
                f"║  Lý do: {err[:50]:<50} ║\n"
                "╠══════════════════════════════════════════════════════════╣\n"
                "║  Chạy KICH_HOAT_BAN_QUYEN.bat để kích hoạt license.    ║\n"
                "╚══════════════════════════════════════════════════════════╝"
            )
        return payload
    except ImportError:
        return {}


# ── 1. DOCX có MathType OLE → LaTeX ─────────────────────────────────────────

def extract_latex_from_mathtype_docx(
    docx_path: str,
    timeout: int = 120,
    wrap_dollar: bool = True,
) -> Optional[str]:
    """
    Upload file DOCX chứa công thức MathType OLE (Equation.DSMT4) lên
    mathtype-latex-api → server dùng MTEF-py giải mã nhị phân OLE → trả về
    cùng DOCX nhưng công thức đã được thay bằng chuỗi LaTeX (bọc `$...$`).

    Trả về: Nội dung văn bản + LaTeX đã trích xuất (str), hoặc None nếu lỗi.
    """
    _require_license()

    if not os.path.isfile(docx_path):
        print(f"[Lỗi] Không tìm thấy file: {docx_path}")
        return None

    print(f"[*] Đang trích xuất MathType OLE → LaTeX từ: {os.path.basename(docx_path)}")
    print(f"    API: {ENDPOINT_DOCX_READER}")

    boundary = uuid.uuid4().hex
    headers = {'Content-Type': f'multipart/form-data; boundary={boundary}'}

    with open(docx_path, 'rb') as f:
        file_bytes = f.read()

    body = bytearray()
    body.extend(f'--{boundary}\r\n'.encode('utf-8'))
    body.extend(f'Content-Disposition: form-data; name="file"; filename="{os.path.basename(docx_path)}"\r\n'.encode('utf-8'))
    body.extend(b'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n')
    body.extend(file_bytes)

    # Tham số tùy chọn
    body.extend(f'\r\n--{boundary}\r\n'.encode('utf-8'))
    body.extend(b'Content-Disposition: form-data; name="wrap_dollar"\r\n\r\n')
    body.extend(b'true' if wrap_dollar else b'false')

    body.extend(f'\r\n--{boundary}--\r\n'.encode('utf-8'))

    req = urllib.request.Request(
        ENDPOINT_DOCX_READER,
        data=bytes(body), headers=headers, method='POST'
    )

    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            if resp.status == 200:
                # Server trả về DOCX đã chuyển đổi → đọc text từ đó
                docx_bytes = resp.read()
                out_path = docx_path.replace('.docx', '_LATEX_EXTRACTED.docx')
                with open(out_path, 'wb') as f_out:
                    f_out.write(docx_bytes)
                print(f"[OK] File đã chuyển đổi lưu tại: {out_path}")

                # Đọc text từ DOCX đã chuyển đổi
                if HAS_DOCX:
                    return _read_text_from_docx(out_path)
                return f"[Đã lưu file chuyển đổi: {out_path}]"
    except Exception as e:
        print(f"[Lỗi] mathtype-latex-api không phản hồi: {e}")

    return None


# ── 2. OCR Ảnh Công Thức → LaTeX (Mistral AI Vision) ────────────────────────

def ocr_image_to_latex(
    image_paths: list,
    timeout: int = 60,
    wrap_dollar: bool = True,
) -> Optional[dict]:
    """
    Upload ảnh chụp đề thi (PNG/JPG/JPEG) lên mathtype-latex-api.
    Server dùng Mistral AI Vision để nhận diện công thức → trả về LaTeX.

    Args:
        image_paths: Danh sách đường dẫn file ảnh (PNG/JPG/JPEG).
        timeout: Thời gian chờ tối đa (giây).
        wrap_dollar: True → bọc kết quả trong $...$

    Trả về: dict {filename: latex_string} hoặc None nếu lỗi.
    """
    _require_license()

    if not image_paths:
        print("[Lỗi] Không có ảnh nào được cung cấp.")
        return None

    print(f"[*] OCR {len(image_paths)} ảnh công thức → LaTeX (Mistral AI Vision)")
    print(f"    API: {ENDPOINT_IMG2LATEX}")

    # Encode ảnh sang base64
    items = []
    for img_path in image_paths:
        if not os.path.isfile(img_path):
            print(f"  [SKIP] Không tìm thấy: {img_path}")
            continue
        with open(img_path, 'rb') as f:
            b64_data = base64.b64encode(f.read()).decode('utf-8')
        items.append({
            "id": os.path.basename(img_path),
            "b64": b64_data
        })

    if not items:
        print("[Lỗi] Không có ảnh hợp lệ nào.")
        return None

    payload = json.dumps({
        "items": items,
        "wrap": wrap_dollar
    }).encode('utf-8')

    headers = {
        'Content-Type': 'application/json',
        'Accept': 'application/json'
    }

    req = urllib.request.Request(
        ENDPOINT_IMG2LATEX,
        data=payload, headers=headers, method='POST'
    )

    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            if resp.status == 200:
                result = json.loads(resp.read().decode('utf-8'))
                print(f"[OK] OCR thành công {len(result)} ảnh.")
                # In tóm tắt kết quả
                for item_id, latex in result.items():
                    preview = str(latex)[:80].replace('\n', ' ')
                    print(f"  [{item_id}] → {preview}...")
                return result
    except Exception as e:
        print(f"[Lỗi] OCR API không phản hồi: {e}")

    return None


# ── 3. Đọc DOCX thường (không MathType) → Text + LaTeX ─────────────────────

def _read_text_from_docx(docx_path: str) -> str:
    """Trích xuất toàn bộ văn bản từ DOCX (python-docx)."""
    if not HAS_DOCX:
        return f"[Lỗi] python-docx chưa được cài đặt. Chạy: pip install python-docx"
    try:
        doc = DocxDocument(docx_path)
        lines = []
        for para in doc.paragraphs:
            if para.text.strip():
                lines.append(para.text)
        # Trích xuất text trong bảng
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    lines.append(row_text)
        return "\n".join(lines)
    except Exception as e:
        return f"[Lỗi khi đọc DOCX: {e}]"


def read_exam_docx(docx_path: str) -> str:
    """
    Đọc file DOCX đề thi thông thường (không có MathType OLE phức tạp).
    Trích xuất toàn bộ văn bản và công thức LaTeX nếu có.
    """
    _require_license()
    print(f"[*] Đọc nội dung đề thi: {os.path.basename(docx_path)}")
    text = _read_text_from_docx(docx_path)
    print(f"[OK] Đã trích xuất {len(text)} ký tự.")
    return text


# ── 4. Đọc PDF → Text ────────────────────────────────────────────────────────

def read_exam_pdf(pdf_path: str) -> str:
    """
    Đọc file PDF đề thi và trích xuất toàn bộ văn bản.
    Dùng PyMuPDF (fitz). Ghi chú: PDF scan (ảnh) sẽ không có text.
    """
    _require_license()

    if not HAS_PYMUPDF:
        return "[Lỗi] PyMuPDF (fitz) chưa cài đặt. Chạy: pip install pymupdf"

    print(f"[*] Đọc nội dung PDF: {os.path.basename(pdf_path)}")
    try:
        doc = fitz.open(pdf_path)
        pages_text = []
        for i, page in enumerate(doc):
            text = page.get_text("text").strip()
            if text:
                pages_text.append(f"--- Trang {i+1} ---\n{text}")
        doc.close()
        full_text = "\n\n".join(pages_text)
        print(f"[OK] Đã trích xuất {len(full_text)} ký tự từ {len(doc)} trang.")
        return full_text
    except Exception as e:
        return f"[Lỗi khi đọc PDF: {e}]"


# ── 5. Hàm tổng hợp: Tự phát hiện loại file ─────────────────────────────────

def read_exam_source(file_path: str, is_mathtype_docx: bool = False) -> str:
    """
    Hàm tổng hợp tự nhận diện loại file đề thi và trích xuất nội dung.

    Args:
        file_path: Đường dẫn tới file đề thi (PDF, DOCX, PNG, JPG, JPEG).
        is_mathtype_docx: True nếu biết DOCX có chứa MathType OLE objects
                          (cần gọi API để chuyển về LaTeX).
    Trả về: Chuỗi văn bản đề thi đầy đủ (kèm LaTeX nếu trích xuất được).
    """
    ext = Path(file_path).suffix.lower()

    if ext == '.pdf':
        return read_exam_pdf(file_path)

    elif ext == '.docx':
        if is_mathtype_docx:
            result = extract_latex_from_mathtype_docx(file_path)
            return result if result else read_exam_docx(file_path)
        else:
            return read_exam_docx(file_path)

    elif ext in ('.png', '.jpg', '.jpeg', '.bmp', '.tiff'):
        result = ocr_image_to_latex([file_path])
        if result:
            parts = []
            for img_id, latex in result.items():
                parts.append(f"[{img_id}]\n{latex}")
            return "\n\n".join(parts)
        return "[OCR thất bại - không có kết quả]"

    else:
        return f"[Lỗi] Định dạng file chưa được hỗ trợ: {ext}"


# ── CLI ───────────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    import argparse

    parser = argparse.ArgumentParser(
        description="Trích xuất nội dung đề thi từ PDF, DOCX (MathType OLE), hoặc ảnh scan"
    )
    parser.add_argument("--file", "-f", required=True,
                        help="Đường dẫn file đề thi (PDF, DOCX, PNG, JPG)")
    parser.add_argument("--mathtype", action="store_true",
                        help="Đánh dấu DOCX có MathType OLE → gọi API chuyển về LaTeX")
    parser.add_argument("--out", "-o", default=None,
                        help="Lưu kết quả văn bản ra file .txt (tùy chọn)")
    args = parser.parse_args()

    content = read_exam_source(args.file, is_mathtype_docx=args.mathtype)

    if args.out:
        with open(args.out, 'w', encoding='utf-8') as f_out:
            f_out.write(content)
        print(f"[OK] Đã lưu kết quả vào: {args.out}")
    else:
        print("\n" + "="*60)
        print("NỘI DUNG TRÍCH XUẤT:")
        print("="*60)
        print(content[:3000])
        if len(content) > 3000:
            print(f"\n... (còn {len(content)-3000} ký tự nữa, dùng --out để lưu toàn bộ)")
