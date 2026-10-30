#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Công Cụ Đọc Nội Dung Slide (Bao gồm Textbox, Bảng biểu & Nhóm hình) & Ghi Lời Bình Vào Phần Notes
Tác giả: Thầy Nguyễn Văn Thiện - THCS Ứng Hòe
- Đọc tiêu đề, nội dung chữ, bảng biểu của từng slide trong file PPTX có sẵn
- Tự động ghi lời giảng/lời bình của giáo viên vào mục Notes dưới chân slide
- Phục vụ việc trình chiếu, lồng tiếng AI và xuất video bài giảng MP4
"""

import sys
import os
import json
import argparse
import subprocess

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    from pptx import Presentation
except ImportError:
    print("[*] Đang tự động cài đặt thư viện python-pptx (lần đầu tiên)...")
    subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx", "--quiet"])
    try:
        from pptx import Presentation
    except ImportError:
        print("[!] Không thể tự động cài đặt python-pptx. Vui lòng chạy 1_CAI_DAT_MOI_TRUONG.bat.")
        sys.exit(1)


def _collect_shape_texts(shape, out_list):
    """Đệ quy trích xuất chữ từ TextFrame, Bảng (Table) và Nhóm hình (GroupShape)."""
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            t = paragraph.text.strip()
            if t:
                out_list.append(t)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip()]
            if cells:
                out_list.append(" | ".join(cells))
    if hasattr(shape, "shapes"):
        for sub_sh in shape.shapes:
            _collect_shape_texts(sub_sh, out_list)


def extract_slide_content(pptx_path):
    """Trích xuất toàn bộ văn bản của từng slide để AI phân tích viết lời bình"""
    pptx_path = os.path.abspath(pptx_path)
    if not os.path.exists(pptx_path):
        print(f"[X] Không tìm thấy file: {pptx_path}")
        return None

    prs = Presentation(pptx_path)
    slides_data = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            _collect_shape_texts(shape, texts)

        existing_notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            existing_notes = slide.notes_slide.notes_text_frame.text.strip()

        slides_data.append(
            {
                "slide_index": idx,
                "title": texts[0] if texts else f"Slide {idx}",
                "content": texts,
                "existing_notes": existing_notes,
            }
        )

    return slides_data


def inject_notes_to_pptx(pptx_path, notes_dict, output_path=None):
    """
    Ghi danh sách lời bình vào mục Notes của từng slide.
    notes_dict: dict {slide_index (int hoặc str): "nội dung lời bình (str)"}
    """
    pptx_path = os.path.abspath(pptx_path)
    if not os.path.exists(pptx_path):
        print(f"[X] Không tìm thấy file: {pptx_path}")
        return False

    if not output_path:
        base, ext = os.path.splitext(pptx_path)
        output_path = f"{base}_CoLoiBinh{ext}"
    else:
        output_path = os.path.abspath(output_path)

    prs = Presentation(pptx_path)
    written_count = 0

    if isinstance(notes_dict, list):
        normalized = {}
        for i, item in enumerate(notes_dict, start=1):
            if isinstance(item, dict):
                s_idx = item.get("slide_index", item.get("slide", item.get("id", i)))
                s_txt = item.get("notes", item.get("text", item.get("script", item.get("existing_notes", ""))))
                normalized[str(s_idx)] = s_txt
            elif isinstance(item, str):
                normalized[str(i)] = item
        notes_dict = normalized

    from copy import deepcopy
    from pptx.util import Inches

    template_notes_sp = None
    for s in prs.slides:
        try:
            ns = s.notes_slide
            if ns.notes_placeholder is not None:
                template_notes_sp = deepcopy(ns.notes_placeholder._element)
                break
        except Exception:
            pass

    for idx, slide in enumerate(prs.slides, start=1):
        str_key = str(idx)
        script_text = notes_dict.get(idx) if idx in notes_dict else notes_dict.get(str_key)
        if script_text and str(script_text).strip():
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            if text_frame is None and template_notes_sp is not None:
                notes_slide.shapes._spTree.append(deepcopy(template_notes_sp))
                text_frame = notes_slide.notes_text_frame
            if text_frame is None:
                tb = notes_slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(4))
                text_frame = tb.text_frame
            text_frame.text = str(script_text).strip()
            written_count += 1

    prs.save(output_path)
    print(f"[OK] Đã ghi thành công lời bình vào phần Notes của {written_count}/{len(prs.slides)} slide!")
    print(f"[OK] File mới đã lưu tại: {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Trích xuất nội dung hoặc ghi lời bình vào PowerPoint Notes")
    parser.add_argument("pptx", help="Đường dẫn file PowerPoint (.pptx)")
    parser.add_argument("--extract", action="store_true", help="Trích xuất nội dung các slide dạng JSON ra màn hình")
    parser.add_argument("--inject-json", type=str, help="Đường dẫn file JSON chứa lời bình để ghi vào Notes")
    parser.add_argument("-o", "--output", help="Đường dẫn file PPTX lưu ra")
    args = parser.parse_args()

    if args.extract:
        data = extract_slide_content(args.pptx)
        if data:
            print(json.dumps(data, ensure_ascii=False, indent=2))
        return

    if args.inject_json:
        if not os.path.exists(args.inject_json):
            print(f"[X] Không tìm thấy file JSON: {args.inject_json}")
            sys.exit(1)
        with open(args.inject_json, "r", encoding="utf-8") as f:
            notes_data = json.load(f)
        inject_notes_to_pptx(args.pptx, notes_data, args.output)
        return

    print("Vui lòng chọn cờ --extract để đọc nội dung slide, hoặc --inject-json để ghi lời bình.")


if __name__ == "__main__":
    main()
