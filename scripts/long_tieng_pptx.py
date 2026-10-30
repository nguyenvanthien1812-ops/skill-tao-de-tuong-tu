#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Hệ Thống Tự Động Lồng Tiếng AI Chuẩn Phát Thanh Viên Cho Slide PowerPoint
Tác giả: Thầy Nguyễn Văn Thiện - THCS Ứng Hòe
- Đọc nội dung lời giảng trong phần Notes của từng slide
- Sử dụng giọng đọc AI tiếng Việt chất lượng cao của Microsoft (Edge-TTS):
  + vi-VN-HoaiMyNeural (Cô Hoài My - Giọng nữ miền Bắc nhẹ nhàng, truyền cảm)
  + vi-VN-NamMinhNeural (Thầy Nam Minh - Giọng nam miền Bắc dõng dạc, trầm ấm)
- Đo chính xác 100% thời lượng từng file âm thanh MP3 (qua mutagen + bộ đọc khung MPEG dự phòng)
- Tạo kịch bản PowerShell lưu chuẩn UTF-8 BOM (utf-8-sig) + chuỗi đơn an toàn tuyệt đối với tên file tiếng Việt có dấu
- Tự động đặt thời gian chuyển slide (AdvanceTime) và xuất bảng mốc giây _SlideTimeline.json phục vụ Video Tương Tác!
"""

import sys
import os
import re
import json
import asyncio
import argparse
import subprocess
import tempfile
import shutil

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    from pptx import Presentation
except ImportError:
    print("[*] Đang tự động cài đặt thư viện python-pptx (lần đầu tiên)...")
    subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx", "--quiet"])
    try:
        from pptx import Presentation
    except ImportError:
        print("[!] Không thể tự động cài đặt python-pptx. Vui lòng chạy 1_CAI_DAT_MOI_TRUONG.bat.")
        sys.exit(1)

try:
    import edge_tts
except ImportError:
    print("[*] Đang tự động cài đặt thư viện giọng đọc AI edge-tts...")
    subprocess.run([sys.executable, "-m", "pip", "install", "edge-tts", "--quiet"])
    import edge_tts


VOICE_MAP = {
    "1": ("vi-VN-HoaiMyNeural", "Cô Hoài My (Nữ miền Bắc - Nhẹ nhàng, truyền cảm)"),
    "2": ("vi-VN-NamMinhNeural", "Thầy Nam Minh (Nam miền Bắc - Dõng dạc, trầm ấm)"),
    "hoaimy": ("vi-VN-HoaiMyNeural", "Cô Hoài My (Nữ miền Bắc)"),
    "namminh": ("vi-VN-NamMinhNeural", "Thầy Nam Minh (Nam miền Bắc)"),
}


def clean_text_for_tts(raw_text: str) -> str:
    """Làm sạch nhẹ các ký tự trang trí Markdown để giọng đọc AI phát âm tự nhiên, không đọc ký hiệu lạ."""
    txt = str(raw_text).strip()
    txt = re.sub(r"[*_`#]+", "", txt)
    txt = txt.replace("->", ", ").replace("→", ", ").replace("⇒", " suy ra ")
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt


def get_exact_mp3_duration(mp3_path: str, fallback_text: str = "") -> float:
    """
    Đo chính xác thời lượng giây của file MP3 vừa sinh bởi Edge-TTS:
    1. Ưu tiên thư viện mutagen.mp3.MP3 (chính xác 100%).
    2. Dự phòng: tự đọc header khung MPEG Layer III thuần Python (Edge-TTS xuất MP3 24kHz mono 48kbps CBR).
    3. Dự phòng cuối: ước lượng theo tốc độ đọc tiếng Việt (~13.5 ký tự/giây).
    """
    try:
        from mutagen.mp3 import MP3
        audio = MP3(mp3_path)
        if audio.info and audio.info.length > 0.2:
            return float(audio.info.length)
    except Exception:
        pass

    # Bộ đọc khung MPEG Audio thuần Python (không cần cài thêm thư viện ngoài)
    try:
        data = open(mp3_path, "rb").read()
        pos = 0
        if data[:3] == b"ID3" and len(data) >= 10:
            tag_size = (
                ((data[6] & 0x7F) << 21)
                | ((data[7] & 0x7F) << 14)
                | ((data[8] & 0x7F) << 7)
                | (data[9] & 0x7F)
            )
            pos = 10 + tag_size

        br_v1_l3 = [0, 32, 40, 48, 56, 64, 80, 96, 112, 128, 160, 192, 224, 256, 320, 0]
        br_v2_l3 = [0, 8, 16, 24, 32, 40, 48, 56, 64, 80, 96, 112, 128, 144, 160, 0]
        sr_table = {
            3: [44100, 48000, 32000, 0],  # MPEG 1
            2: [22050, 24000, 16000, 0],  # MPEG 2
            0: [11025, 12000, 8000, 0],   # MPEG 2.5
        }
        total_sec = 0.0
        frames = 0
        n = len(data)
        while pos + 4 <= n:
            if data[pos] == 0xFF and (data[pos + 1] & 0xE0) == 0xE0:
                b1, b2 = data[pos + 1], data[pos + 2]
                ver_id = (b1 >> 3) & 0x03
                layer = (b1 >> 1) & 0x03
                br_idx = (b2 >> 4) & 0x0F
                sr_idx = (b2 >> 2) & 0x03
                pad = (b2 >> 1) & 0x01
                if ver_id != 1 and layer == 1 and 0 < br_idx < 15 and sr_idx < 3:
                    sr = sr_table[ver_id][sr_idx]
                    br = (br_v1_l3[br_idx] if ver_id == 3 else br_v2_l3[br_idx]) * 1000
                    samples = 1152 if ver_id == 3 else 576
                    frame_len = int((samples * br // 8) // sr + pad)
                    if frame_len > 0:
                        total_sec += samples / float(sr)
                        pos += frame_len
                        frames += 1
                        continue
            pos += 1
        if frames > 5 and total_sec > 0.2:
            return total_sec
    except Exception:
        pass

    est = max(4.0, len(fallback_text) / 13.5)
    return est


def ps_quote(val: str) -> str:
    """Đặt đường dẫn vào cặp nháy đơn PowerShell an toàn tuyệt đối với tiếng Việt, dấu cách, ký tự $."""
    return "'" + str(val).replace("'", "''") + "'"


async def generate_speech(text, voice_name, output_path):
    """Sinh file âm thanh mp3 từ văn bản qua edge-tts (tự thử lại tối đa 3 lần nếu mạng chập chờn)."""
    last_err = None
    for attempt in range(3):
        try:
            communicate = edge_tts.Communicate(text, voice_name, rate="+0%", volume="+0%")
            await communicate.save(output_path)
            if os.path.exists(output_path) and os.path.getsize(output_path) > 0:
                return
        except Exception as e:
            last_err = e
            await asyncio.sleep(1.0 * (attempt + 1))
    if last_err:
        raise last_err


def process_long_tieng(pptx_path, voice_choice="1", output_pptx=None):
    pptx_path = os.path.abspath(pptx_path)
    if not os.path.exists(pptx_path):
        print(f"[X] Không tìm thấy file: {pptx_path}")
        return False

    voice_key = str(voice_choice).lower().strip()
    voice_info = VOICE_MAP.get(voice_key, VOICE_MAP["1"])
    voice_id, voice_label = voice_info

    if not output_pptx:
        base, ext = os.path.splitext(pptx_path)
        output_pptx = f"{base}_DaLongTieng{ext}"
    else:
        output_pptx = os.path.abspath(output_pptx)

    print("=" * 68)
    print("      HỆ THỐNG TỰ ĐỘNG LỒNG TIẾNG AI CHO SLIDE POWERPOINT")
    print("        Tác giả: Thầy Nguyễn Văn Thiện - THCS Ứng Hòe")
    print("=" * 68)
    print(f"\n[*] File slide gốc    : {os.path.basename(pptx_path)}")
    print(f"[*] Giọng đọc AI chọn : {voice_label}")

    prs = Presentation(pptx_path)
    total_slides = len(prs.slides)
    print(f"[*] Tổng số slide     : {total_slides} trang")

    temp_dir = tempfile.mkdtemp(prefix="audio_slides_")
    audio_files = {}
    py_durations = {}

    print("\n[*] Đang chuyển lời giảng (Notes) thành giọng đọc AI:")

    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)

    has_any_notes = False
    for idx, slide in enumerate(prs.slides, start=1):
        note_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note_text = clean_text_for_tts(slide.notes_slide.notes_text_frame.text)

        if note_text:
            has_any_notes = True
            mp3_path = os.path.join(temp_dir, f"slide_{idx:02d}.mp3")
            print(f"    • Slide {idx:02d}: Đang đọc ({len(note_text)} ký tự)...", end="", flush=True)
            try:
                loop.run_until_complete(generate_speech(note_text, voice_id, mp3_path))
                if os.path.exists(mp3_path) and os.path.getsize(mp3_path) > 0:
                    audio_files[idx] = mp3_path
                    exact_len = get_exact_mp3_duration(mp3_path, note_text)
                    # Cộng thêm 1.0 giây nghỉ tự nhiên giữa 2 slide
                    slide_dur = round(exact_len + 1.0, 1)
                    py_durations[idx] = slide_dur
                    print(f" [OK - {slide_dur:.1f}s]")
                else:
                    print(" [Trống]")
            except Exception as e:
                print(f"\n      [!] Lỗi khi đọc slide {idx}: {e}")
        else:
            py_durations[idx] = 5.0
            print(f"    • Slide {idx:02d}: Không có lời giảng Notes (đặt mặc định 5.0s)")

    loop.close()

    if not has_any_notes or not audio_files:
        print("\n[!] Chú ý: File PowerPoint này chưa có lời bình trong phần Notes của các slide.")
        print("    Vui lòng nhờ Trợ lý AI viết lời bình vào Notes trước (hoặc dùng them_loi_binh_pptx.py).")
        shutil.rmtree(temp_dir, ignore_errors=True)
        return False

    print(f"\n[*] Đã sinh xong {len(audio_files)} đoạn giọng đọc AI.")
    print("[*] Đang tự động nhúng âm thanh và cài thời gian chuyển trang vào PowerPoint...")

    timeline_json_path = os.path.splitext(output_pptx)[0] + "_SlideTimeline.json"
    timeline_raw_path = os.path.join(temp_dir, "slide_durations.txt")

    # Tạo script PowerShell chuẩn UTF-8 BOM (utf-8-sig) để PowerShell 5.1 đọc chuẩn 100% đường dẫn tiếng Việt có dấu
    ps_lines = [
        '$ErrorActionPreference = "Stop"',
        f"$pptxPath = {ps_quote(pptx_path)}",
        f"$outPath = {ps_quote(output_pptx)}",
        f"$durLogPath = {ps_quote(timeline_raw_path)}",
        "if (Test-Path -LiteralPath $outPath) { Remove-Item -LiteralPath $outPath -Force }",
        "$durLines = @()",
        "$createdNewPP = $false",
        "$powerPoint = $null",
        "$presentation = $null",
        "try {",
        "    try {",
        '        $powerPoint = [Runtime.InteropServices.Marshal]::GetActiveObject("PowerPoint.Application")',
        "    } catch {",
        "        $powerPoint = New-Object -ComObject PowerPoint.Application",
        "        $createdNewPP = $true",
        "    }",
        "    $presentation = $powerPoint.Presentations.Open($pptxPath, [Microsoft.Office.Core.MsoTriState]::msoTrue, [Microsoft.Office.Core.MsoTriState]::msoFalse, [Microsoft.Office.Core.MsoTriState]::msoFalse)",
    ]

    for idx in range(1, total_slides + 1):
        mp3 = audio_files.get(idx)
        py_dur = py_durations.get(idx, 6.0)
        if mp3:
            ps_lines.append(
                f"""
    try {{
        $slide = $presentation.Slides.Item({idx})
        $shape = $slide.Shapes.AddMediaObject2({ps_quote(mp3)}, $false, $true, 20, 20, 40, 40)
        $shape.AnimationSettings.PlaySettings.PlayOnEntry = [Microsoft.Office.Core.MsoTriState]::msoTrue
        $shape.AnimationSettings.PlaySettings.HideWhileNotPlaying = [Microsoft.Office.Core.MsoTriState]::msoTrue
        $slide.SlideShowTransition.AdvanceOnTime = [Microsoft.Office.Core.MsoTriState]::msoTrue
        $durSec = {py_dur}
        $lenMs = 0
        try {{ $lenMs = [int]$shape.MediaFormat.Length }} catch {{ $lenMs = 0 }}
        if ($lenMs -gt 200) {{
            $comSec = [math]::Round(($lenMs / 1000.0) + 1.0, 1)
            if ($comSec -gt $durSec) {{ $durSec = $comSec }}
        }}
        $slide.SlideShowTransition.AdvanceTime = $durSec
        $durLines += "{idx}|$durSec"
    }} catch {{
        $durLines += "{idx}|{py_dur}"
    }}"""
            )
        else:
            ps_lines.append(
                f"""
    try {{
        $slide = $presentation.Slides.Item({idx})
        $slide.SlideShowTransition.AdvanceOnTime = [Microsoft.Office.Core.MsoTriState]::msoTrue
        $slide.SlideShowTransition.AdvanceTime = 5.0
    }} catch {{}}
    $durLines += "{idx}|5.0" """
            )

    ps_lines.extend(
        [
            "    $durLines | Set-Content -LiteralPath $durLogPath -Encoding UTF8",
            "    $presentation.SaveAs($outPath)",
            '    Write-Host "[OK] Da nhung thanh cong toan bo am thanh!"',
            "} finally {",
            "    if ($presentation -ne $null) { try { $presentation.Close() } catch {} }",
            "    if ($createdNewPP -and $powerPoint -ne $null) { try { $powerPoint.Quit() } catch {} }",
            "}",
        ]
    )

    ps_script_path = os.path.join(temp_dir, "embed_audio.ps1")
    # CỰC KỲ QUAN TRỌNG: Phải ghi utf-8-sig (có BOM) để Windows PowerShell 5.1 đọc đúng tên file tiếng Việt!
    with open(ps_script_path, "w", encoding="utf-8-sig") as f:
        f.write("\n".join(ps_lines))

    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", ps_script_path],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=240,
        )
        if os.path.exists(output_pptx) and os.path.getsize(output_pptx) > 0:
            try:
                dur_map = dict(py_durations)
                if os.path.exists(timeline_raw_path):
                    with open(timeline_raw_path, "r", encoding="utf-8-sig") as df:
                        for line in df:
                            parts = line.strip().split("|")
                            if len(parts) == 2:
                                dur_map[int(parts[0])] = float(parts[1])
                cum_sec = 0.0
                timeline_entries = []
                for idx, slide in enumerate(prs.slides, start=1):
                    d_sec = dur_map.get(idx, 5.0)
                    start_s = round(cum_sec, 1)
                    cum_sec += d_sec
                    end_s = round(cum_sec, 1)
                    titles = [
                        p.text.strip()
                        for sh in slide.shapes
                        if getattr(sh, "has_text_frame", False)
                        for p in sh.text_frame.paragraphs
                        if p.text.strip()
                    ]
                    mm, ss = divmod(int(end_s), 60)
                    timeline_entries.append(
                        {
                            "slide": idx,
                            "title": titles[0] if titles else f"Slide {idx}",
                            "start_sec": start_s,
                            "duration_sec": d_sec,
                            "end_sec": int(end_s),
                            "end_mmss": f"{mm:02d}:{ss:02d}",
                        }
                    )
                with open(timeline_json_path, "w", encoding="utf-8") as tf:
                    json.dump(
                        {
                            "total_duration_sec": int(round(cum_sec)),
                            "slides": timeline_entries,
                        },
                        tf,
                        ensure_ascii=False,
                        indent=2,
                    )
                print(f"  • Bảng mốc giây từng Slide (để gắn câu hỏi Video Tương Tác): {timeline_json_path}")
            except Exception as te:
                print(f"  [!] Ghi chú: Không xuất được file timeline phụ ({te})")

            print("\n" + "=" * 68)
            print("  🎉 [LỒNG TIẾNG AI THÀNH CÔNG!]")
            print(f"  • File PowerPoint đã lồng tiếng: {output_pptx}")
            print("  • Mỗi slide đã có sẵn giọng đọc AI phát tự động & khớp thời gian chuyển trang!")
            print("  • Bước tiếp: Kéo file này vào 'XUAT_VIDEO_TU_PPTX.bat' để xuất Video MP4 Full HD!")
            print("=" * 68)
            return True
        else:
            print(f"[!] Lỗi PowerPoint COM: {proc.stderr.strip() or proc.stdout.strip()}")
            return False
    except Exception as e:
        print(f"[!] Lỗi khi nhúng âm thanh: {e}")
        return False
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def main():
    parser = argparse.ArgumentParser(description="Lồng tiếng AI tự động vào slide PowerPoint")
    parser.add_argument("pptx", help="Đường dẫn file PowerPoint (.pptx)")
    parser.add_argument("--voice", default="1", help="1: Cô Hoài My, 2: Thầy Nam Minh")
    parser.add_argument("-o", "--output", help="Đường dẫn file xuất ra")
    args = parser.parse_args()

    ok = process_long_tieng(args.pptx, args.voice, args.output)
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
