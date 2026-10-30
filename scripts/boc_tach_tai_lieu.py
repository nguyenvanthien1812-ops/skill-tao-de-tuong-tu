#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Công Cụ Tự Động Bóc Tách Văn Bản & Hình Ảnh Từ File Giáo Án Word (DOCX) & SGK (PDF)
Tác giả: Thầy Nguyễn Văn Thiện - THCS Ứng Hòe
- Bóc tách toàn bộ hình ảnh phân giải gốc từ file Word (.docx) hoặc Sách giáo khoa (.pdf)
- Trích xuất toàn bộ văn bản để AI phân tích cấu trúc bài học và lập Storyboard
"""

import os
import sys
import json
import zipfile
import argparse
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass


def extract_docx(docx_path, out_dir):
    """Bóc tách ảnh và văn bản từ file Word (.docx)"""
    docx_path = Path(docx_path).resolve()
    out_dir = Path(out_dir).resolve()
    images_dir = out_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    extracted_images = []
    text_content = []

    # 1. Trích xuất hình ảnh từ gói media của DOCX
    try:
        with zipfile.ZipFile(docx_path, 'r') as z:
            media_files = [f for f in z.namelist() if f.startswith('word/media/')]
            for idx, mf in enumerate(media_files, start=1):
                ext = Path(mf).suffix
                out_name = f"image_{idx:02d}{ext}"
                out_path = images_dir / out_name
                with open(out_path, 'wb') as f_out:
                    f_out.write(z.read(mf))
                extracted_images.append({
                    "name": out_name,
                    "path": str(out_path),
                    "size_bytes": out_path.stat().st_size
                })
    except Exception as e:
        print(f"[!] Canh bao khi trich xuat anh tu DOCX: {e}")

    # 2. Trích xuất văn bản từ DOCX
    try:
        import docx
        doc = docx.Document(str(docx_path))
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                text_content.append(t)
        for tbl in doc.tables:
            for row in tbl.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_content.append(" | ".join(row_texts))
    except ImportError:
        # Fallback đọc thẳng word/document.xml nếu máy chưa có python-docx
        try:
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(docx_path, 'r') as z:
                xml_data = z.read('word/document.xml')
                root = ET.fromstring(xml_data)
                for elem in root.iter():
                    if elem.tag.endswith('}t') and elem.text:
                        text_content.append(elem.text.strip())
        except Exception as e:
            print(f"[!] Khong the doc van ban docx: {e}")

    # Ghi văn bản ra file txt
    txt_path = out_dir / "content.txt"
    txt_path.write_text("\n\n".join(text_content), encoding='utf-8')

    return {
        "source": str(docx_path),
        "total_paragraphs": len(text_content),
        "total_images": len(extracted_images),
        "images": extracted_images,
        "content_txt": str(txt_path)
    }


def extract_pdf(pdf_path, out_dir):
    """Bóc tách ảnh và văn bản từ file Sách giáo khoa PDF"""
    pdf_path = Path(pdf_path).resolve()
    out_dir = Path(out_dir).resolve()
    images_dir = out_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    extracted_images = []
    text_content = []

    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(pdf_path))
        img_idx = 1
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            text = page.get_text()
            if text.strip():
                text_content.append(f"--- TRANG {page_idx + 1} ---\n" + text.strip())

            # Bóc tách ảnh nhúng trong trang
            image_list = page.get_images(full=True)
            for img_info in image_list:
                xref = img_info[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                image_ext = base_image["ext"]
                out_name = f"page{page_idx + 1:02d}_img{img_idx:02d}.{image_ext}"
                out_path = images_dir / out_name
                out_path.write_bytes(image_bytes)
                extracted_images.append({
                    "name": out_name,
                    "path": str(out_path),
                    "page": page_idx + 1,
                    "size_bytes": len(image_bytes)
                })
                img_idx += 1
        doc.close()
    except ImportError:
        # Fallback dung pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(str(pdf_path))
            img_idx = 1
            for p_idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_content.append(f"--- TRANG {p_idx + 1} ---\n" + text.strip())
                for count, image_file_object in enumerate(page.images):
                    out_name = f"page{p_idx + 1:02d}_img{img_idx:02d}_{image_file_object.name}"
                    out_path = images_dir / out_name
                    out_path.write_bytes(image_file_object.data)
                    extracted_images.append({
                        "name": out_name,
                        "path": str(out_path),
                        "page": p_idx + 1,
                        "size_bytes": len(image_file_object.data)
                    })
                    img_idx += 1
        except Exception as e:
            print(f"[!] Loi khi doc PDF: {e}")

    # Ghi văn bản ra file txt
    txt_path = out_dir / "content.txt"
    txt_path.write_text("\n\n".join(text_content), encoding='utf-8')

    return {
        "source": str(pdf_path),
        "total_paragraphs": len(text_content),
        "total_images": len(extracted_images),
        "images": extracted_images,
        "content_txt": str(txt_path)
    }


def main():
    parser = argparse.ArgumentParser(description="Cong cu boc tach tai lieu Word & PDF")
    parser.add_argument("input_file", help="Duong dan den file DOCX hoac PDF")
    parser.add_argument("-o", "--output", default="extracted_assets", help="Thu muc luu ket qua")
    args = parser.parse_args()

    in_file = Path(args.input_file)
    if not in_file.exists():
        print(f"[X] Khong tim thay file: {in_file}")
        sys.exit(1)

    suffix = in_file.suffix.lower()
    if suffix == ".docx":
        res = extract_docx(in_file, args.output)
    elif suffix == ".pdf":
        res = extract_pdf(in_file, args.output)
    else:
        print(f"[X] Dinh dang file khong ho tro: {suffix}. Chi ho tro .docx hoac .pdf")
        sys.exit(1)

    print("=" * 65)
    print("      🎉 BÓC TÁCH NỘI DUNG TÀI LIỆU THÀNH CÔNG!")
    print("=" * 65)
    print(f"[*] File nguồn          : {res['source']}")
    print(f"[*] Số đoạn văn bản     : {res['total_paragraphs']}")
    print(f"[*] Số hình ảnh bóc tách: {res['total_images']}")
    print(f"[*] Thư mục lưu ảnh     : {Path(args.output).resolve() / 'images'}")
    print(f"[*] File text trích xuất: {res['content_txt']}")
    print("=" * 65)


if __name__ == "__main__":
    main()
