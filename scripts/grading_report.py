#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
grading_report.py — Xuất Báo Cáo Chấm Bài Tự Luận
====================================================
Xuất kết quả chấm từ AIGrader sang Word / PDF / Excel / Chat summary.
"""

import os
try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
except ImportError:
    pass

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
except ImportError:
    pass

class GradingReport:
    def __init__(self, results: list, subject: str, exam_code: str = ''):
        self.results = results
        self.subject = subject
        self.exam_code = exam_code

    def to_chat_summary(self, student_name: str = 'Học sinh') -> str:
        if not self.results:
            return "Chưa có kết quả chấm bài."
        
        res = self.results[0]
        score = res.get('total_score', 0)
        max_score = res.get('max_score', 10)
        
        emoji_status = {
            'full': '✅',
            'partial': '⚠️',
            'none': '❌'
        }
        
        md = f"## Báo Cáo Chấm Bài: {student_name}\n"
        md += f"**Môn:** {self.subject} | **Điểm AI gợi ý:** {score}/{max_score} đ\n\n"
        md += "### Chi tiết điểm từng phần:\n"
        
        for b in res.get('breakdown', []):
            st = emoji_status.get(b.get('status', 'none'), '➖')
            md += f"- {st} **{b.get('criterion')}** ({b.get('score')}/{b.get('max_score')}đ)\n"
            md += f"  - *Nhận xét:* {b.get('comment')}\n"
            
        md += f"\n### Nhận xét chung:\n{res.get('overall_comment')}\n"
        
        if res.get('strengths'):
            md += "\n**Điểm mạnh:**\n" + "\n".join(f"- {s}" for s in res.get('strengths'))
            
        if res.get('improvement_suggestions'):
            md += "\n\n**Gợi ý cải thiện:**\n" + "\n".join(f"- {i}" for i in res.get('improvement_suggestions'))
            
        md += "\n\n*Thầy/Cô có muốn điều chỉnh điểm nào không?*"
        return md

    def to_excel(self, class_results: list, output_path: str) -> str:
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Kết quả chấm"
            
            headers = ["STT", "Họ tên", "Mã HS", "Điểm gợi ý AI", "Điểm GV xác nhận", "Nhận xét"]
            ws.append(headers)
            
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            
            ws.freeze_panes = "A2"
            
            for idx, r in enumerate(class_results, 1):
                name = r.get("name", f"Học sinh {idx}")
                code = r.get("code", f"HS{idx:03d}")
                score = r.get("result", {}).get("total_score", 0)
                cmt = r.get("result", {}).get("overall_comment", "")
                ws.append([idx, name, code, score, "", cmt])
                
            wb.save(output_path)
            return output_path
        except Exception as e:
            return f"Lỗi xuất Excel: {e}"

    def to_word_feedback(self, student_name: str, output_path: str) -> str:
        try:
            doc = Document()
            
            heading = doc.add_heading(f"PHIẾU NHẬN XÉT BÀI LÀM - {self.subject.upper()}", 0)
            
            doc.add_paragraph(f"Học sinh: {student_name}")
            doc.add_paragraph(f"Mã đề: {self.exam_code}")
            
            res = self.results[0] if self.results else {}
            score = res.get('total_score', 0)
            max_score = res.get('max_score', 10)
            doc.add_paragraph(f"Điểm bài làm: {score}/{max_score}")
            
            doc.add_heading("Chi tiết đánh giá:", level=1)
            
            breakdown = res.get('breakdown', [])
            if breakdown:
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Tiêu chí'
                hdr_cells[1].text = 'Điểm'
                hdr_cells[2].text = 'Nhận xét'
                
                for b in breakdown:
                    row_cells = table.add_row().cells
                    row_cells[0].text = b.get('criterion', '')
                    row_cells[1].text = f"{b.get('score', 0)}/{b.get('max_score', 0)}"
                    row_cells[2].text = b.get('comment', '')
            
            doc.add_heading("Nhận xét chung:", level=1)
            doc.add_paragraph(res.get('overall_comment', ''))
            
            doc.add_paragraph("\n\nChữ ký Giáo viên:\n............................")
            
            doc.save(output_path)
            return output_path
        except Exception as e:
            return f"Lỗi xuất Word: {e}"

    def _gen_feedback_html(self, result: dict, student_name: str) -> str:
        score = result.get('total_score', 0)
        max_score = result.get('max_score', 10)
        
        html = f"""
        <html>
        <head>
            <style>
                body {{ font-family: Arial, sans-serif; margin: 40px; }}
                .header {{ background-color: teal; color: white; padding: 20px; text-align: center; }}
                .score {{ font-size: 24px; font-weight: bold; color: teal; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
                th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
                th {{ background-color: #f2f2f2; }}
                .full {{ background-color: #d4edda; }}
                .partial {{ background-color: #fff3cd; }}
                .none {{ background-color: #f8d7da; }}
            </style>
        </head>
        <body>
            <div class="header">
                <h1>Phiếu Nhận Xét Bài Làm</h1>
                <h2>{student_name} - Môn: {self.subject}</h2>
            </div>
            <p class="score">Điểm: {score}/{max_score}</p>
            <h3>Chi tiết:</h3>
            <table>
                <tr>
                    <th>Tiêu chí</th>
                    <th>Điểm</th>
                    <th>Nhận xét</th>
                </tr>
        """
        for b in result.get('breakdown', []):
            st = b.get('status', 'none')
            html += f"""
                <tr class="{st}">
                    <td>{b.get('criterion')}</td>
                    <td>{b.get('score')}/{b.get('max_score')}</td>
                    <td>{b.get('comment')}</td>
                </tr>
            """
        html += f"""
            </table>
            <h3>Nhận xét chung:</h3>
            <p>{result.get('overall_comment', '')}</p>
        </body>
        </html>
        """
        return html

    def to_pdf(self, student_name: str, output_path: str) -> str:
        # Mocking PDF generation
        try:
            res = self.results[0] if self.results else {}
            html = self._gen_feedback_html(res, student_name)
            with open(output_path.replace('.pdf', '.html'), 'w', encoding='utf-8') as f:
                f.write(html)
            # In a real scenario, use playwright/puppeteer to render HTML to PDF
            return f"{output_path} (Actually saved as HTML for demo)"
        except Exception as e:
            return f"Lỗi xuất PDF: {e}"

    def display_confirmation_menu(self, result: dict) -> str:
        score = result.get('total_score', 0)
        max_score = result.get('max_score', 10)
        return f"Điểm AI gợi ý: {score}/{max_score}. Nhấn Enter đồng ý, hoặc nhập điểm khác:"

if __name__ == '__main__':
    mock_res = [{
        "total_score": 3.5,
        "max_score": 5.0,
        "breakdown": [
            {"criterion": "Mô hình", "score": 0.5, "max_score": 0.5, "status": "full", "comment": "Tốt"}
        ],
        "overall_comment": "Làm bài tốt.",
        "strengths": ["Rõ ràng"],
        "improvement_suggestions": ["Cần cẩn thận hơn"]
    }]
    report = GradingReport(mock_res, "math")
    print(report.to_chat_summary("Nguyễn Văn A"))
