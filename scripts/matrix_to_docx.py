import sys, os, json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "references" / "default_matrices.json"

def _set_cell_style(cell, text, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=None, italic=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    
    if bg_color:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), bg_color)
        tcPr.append(shd)

def _set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def export_matrix_to_docx(matrix_key, output_path=None):
    if not DB_PATH.exists():
        raise FileNotFoundError(f"Không tìm thấy: {DB_PATH}")
    
    data = json.loads(DB_PATH.read_text(encoding='utf-8'))
    if matrix_key not in data:
        raise KeyError(f"Không tìm thấy ma trận: {matrix_key}")
        
    m = data[matrix_key]
    title = m.get("display_name", matrix_key)
    subj = m.get("subject", "Môn học").capitalize()
    grade = m.get("grade", 12)
    duration = m.get("duration_minutes", 90)
    cells = m.get("cells", [])
    
    doc = Document()
    
    # Page Setup: A4 Landscape for standard wide matrix table
    section = doc.sections[0]
    section.orientation = 1  # Landscape
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.5)
    
    # Header School/Department
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_after = Pt(2)
    r1 = p_head.add_run("SỞ GD&ĐT ........................................\nTRƯỜNG THPT ..................................")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.bold = True
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run(f"KHUNG MA TRẬN VÀ BẢN ĐẶC TẢ ĐỀ KIỂM TRA ĐỊNH KỲ\n")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(14)
    r_t.bold = True
    
    r_sub = p_title.add_run(f"MÔN: {subj.upper()} {grade} — THỜI GIAN LÀM BÀI: {duration} PHÚT (CHUẨN GDPT 2018)")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.bold = True
    
    # -------------------------------------------------------------
    # BẢNG 1: KHUNG MA TRẬN
    # -------------------------------------------------------------
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(8)
    p_sec1.paragraph_format.space_after = Pt(4)
    r_s1 = p_sec1.add_run("I. KHUNG MA TRẬN ĐỀ KIỂM TRA (ĐỊNH DẠNG MỚI BỘ GD&ĐT)")
    r_s1.font.name = 'Times New Roman'
    r_s1.font.size = Pt(12)
    r_s1.bold = True
    
    # Table Matrix
    t1 = doc.add_table(rows=2, cols=9)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t1)
    
    # Header row 1
    h_row0 = t1.rows[0].cells
    _set_cell_style(h_row0[0], "TT", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[1], "Chủ đề / Đơn vị kiến thức", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[2], "Dạng thức câu hỏi", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[3], "Nhận biết\n(NB)", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[4], "Thông hiểu\n(TH)", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[5], "Vận dụng\n(VD)", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[6], "Vận dụng cao\n(VDC)", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[7], "Tổng số\ncâu/ý", bold=True, bg_color="D6E4F0")
    _set_cell_style(h_row0[8], "Tổng điểm\n(Dự kiến)", bold=True, bg_color="D6E4F0")
    
    # Header row 2 (Tỉ lệ % điểm)
    h_row1 = t1.rows[1].cells
    _set_cell_style(h_row1[0], "", bg_color="EBF5FB")
    _set_cell_style(h_row1[1], "Tỉ lệ phần trăm điểm định hướng", bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color="EBF5FB")
    _set_cell_style(h_row1[2], "100%", bold=True, bg_color="EBF5FB")
    _set_cell_style(h_row1[3], "30% - 40%", bold=True, bg_color="EBF5FB")
    _set_cell_style(h_row1[4], "30% - 40%", bold=True, bg_color="EBF5FB")
    _set_cell_style(h_row1[5], "20%", bold=True, bg_color="EBF5FB")
    _set_cell_style(h_row1[6], "10%", bold=True, bg_color="EBF5FB")
    _set_cell_style(h_row1[7], "100%", bold=True, bg_color="EBF5FB")
    _set_cell_style(h_row1[8], "10.0 đ", bold=True, bg_color="EBF5FB")
    
    tot_nb = tot_th = tot_vd = tot_vdc = 0
    for idx, cell in enumerate(cells, 1):
        nb = cell.get("nb", 0)
        th = cell.get("th", 0)
        vd = cell.get("vd", 0)
        vdc = cell.get("vdc", 0)
        tot_nb += nb
        tot_th += th
        tot_vd += vd
        tot_vdc += vdc
        sub_tot = nb + th + vd + vdc
        
        part_str = "Phần I (TN 4 lựa chọn)" if cell.get("part") == "part1_mc" else \
                   "Phần II (TN Đúng/Sai 4 ý)" if cell.get("part") == "part2_tf" else \
                   "Phần III (Trả lời ngắn)" if cell.get("part") == "part3_sa" else "Tự luận"
                   
        r = t1.add_row().cells
        _set_cell_style(r[0], str(idx), font_size=10.5)
        _set_cell_style(r[1], cell.get("topic_display", ""), align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10.5)
        _set_cell_style(r[2], part_str, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10)
        _set_cell_style(r[3], str(nb) if nb > 0 else "-", font_size=10.5)
        _set_cell_style(r[4], str(th) if th > 0 else "-", font_size=10.5)
        _set_cell_style(r[5], str(vd) if vd > 0 else "-", font_size=10.5)
        _set_cell_style(r[6], str(vdc) if vdc > 0 else "-", font_size=10.5)
        _set_cell_style(r[7], str(sub_tot), bold=True, font_size=10.5)
        _set_cell_style(r[8], f"{sub_tot * 0.25:.2f}".rstrip('0').rstrip('.') + " đ" if cell.get("part") == "part1_mc" else "-", font_size=10)
        
    # Total row
    r_sum = t1.add_row().cells
    _set_cell_style(r_sum[0], "", bg_color="EAEDED")
    _set_cell_style(r_sum[1], "TỔNG CỘNG TOÀN ĐỀ", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color="EAEDED")
    _set_cell_style(r_sum[2], "Phần I + II + III", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color="EAEDED")
    _set_cell_style(r_sum[3], str(tot_nb), bold=True, bg_color="EAEDED")
    _set_cell_style(r_sum[4], str(tot_th), bold=True, bg_color="EAEDED")
    _set_cell_style(r_sum[5], str(tot_vd), bold=True, bg_color="EAEDED")
    _set_cell_style(r_sum[6], str(tot_vdc), bold=True, bg_color="EAEDED")
    _set_cell_style(r_sum[7], str(tot_nb + tot_th + tot_vd + tot_vdc), bold=True, bg_color="EAEDED")
    _set_cell_style(r_sum[8], "10,0 điểm", bold=True, bg_color="EAEDED")
    
    # -------------------------------------------------------------
    # BẢNG 2: BẢN ĐẶC TẢ CHI TIẾT
    # -------------------------------------------------------------
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(14)
    p_sec2.paragraph_format.space_after = Pt(4)
    r_s2 = p_sec2.add_run("II. BẢN ĐẶC TẢ MA TRẬN ĐỀ KIỂM TRA ĐỊNH KỲ (YÊU CẦU CẦN ĐẠT)")
    r_s2.font.name = 'Times New Roman'
    r_s2.font.size = Pt(12)
    r_s2.bold = True
    
    t2 = doc.add_table(rows=1, cols=7)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t2)
    
    h2 = t2.rows[0].cells
    _set_cell_style(h2[0], "TT", bold=True, bg_color="D5F5E3")
    _set_cell_style(h2[1], "Chủ đề / Bài học", bold=True, bg_color="D5F5E3")
    _set_cell_style(h2[2], "Đơn vị kiến thức", bold=True, bg_color="D5F5E3")
    _set_cell_style(h2[3], "Yêu cầu cần đạt (YCCĐ chuẩn CT GDPT 2018)", bold=True, bg_color="D5F5E3")
    _set_cell_style(h2[4], "Dạng thức", bold=True, bg_color="D5F5E3")
    _set_cell_style(h2[5], "Mức độ", bold=True, bg_color="D5F5E3")
    _set_cell_style(h2[6], "Số câu", bold=True, bg_color="D5F5E3")
    
    for idx, cell in enumerate(cells, 1):
        nb = cell.get("nb", 0)
        th = cell.get("th", 0)
        vd = cell.get("vd", 0)
        vdc = cell.get("vdc", 0)
        part_str = "Phần I" if cell.get("part") == "part1_mc" else \
                   "Phần II" if cell.get("part") == "part2_tf" else \
                   "Phần III" if cell.get("part") == "part3_sa" else "Tự luận"
        
        yccd = f"Nhận biết và vận dụng kiến thức về {cell.get('topic_display', '')}; giải quyết bài toán và tình huống thực tiễn gắn với định hướng GDPT 2018."
        
        levels = []
        if nb: levels.append(f"{nb} NB")
        if th: levels.append(f"{th} TH")
        if vd: levels.append(f"{vd} VD")
        if vdc: levels.append(f"{vdc} VDC")
        lvl_str = ", ".join(levels)
        
        r2 = t2.add_row().cells
        _set_cell_style(r2[0], str(idx), font_size=10)
        _set_cell_style(r2[1], cell.get("topic_display", ""), align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=10)
        _set_cell_style(r2[2], cell.get("topic_display", ""), align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10)
        _set_cell_style(r2[3], yccd, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10)
        _set_cell_style(r2[4], part_str, font_size=10)
        _set_cell_style(r2[5], lvl_str, font_size=10)
        _set_cell_style(r2[6], str(nb + th + vd + vdc), bold=True, font_size=10)
        
    if not output_path:
        output_path = BASE_DIR / f"{matrix_key.upper()}_MA_TRAN_DAC_TA_CHUAN_BGD.docx"
    else:
        output_path = Path(output_path)
        
    doc.save(str(output_path))
    print(f"[OK] Đã xuất file Word Ma trận & Bản đặc tả: {output_path}")
    return output_path

if __name__ == "__main__":
    test_key = sys.argv[1] if len(sys.argv) > 1 else "math_12_thpt_90min"
    export_matrix_to_docx(test_key)
